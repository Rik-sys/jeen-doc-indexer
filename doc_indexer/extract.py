"""Text extraction and cleaning for PDF and DOCX files.

The file type is checked twice: by extension and by the file's magic bytes,
so a renamed file (``notes.txt`` -> ``notes.pdf``) is rejected with a clear
message instead of crashing the parser.
"""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .errors import CorruptedFileError, MissingFileError, NoTextFoundError, UnsupportedFileError

SUPPORTED_EXTENSIONS = (".pdf", ".docx")
PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str  # base name only, e.g. "example.pdf" (never the full local path)
    text: str
    pages: int | None  # PDF page count; None for DOCX

    @property
    def char_count(self) -> int:
        return len(self.text)


# ─────────────────────────── public API ───────────────────────────

def extract_document(path: str | Path) -> ExtractedDocument:
    """Validate ``path``, extract its text and return it cleaned."""
    p = Path(path).expanduser()
    if not p.exists():
        raise MissingFileError(f"File not found: {path}")
    if not p.is_file():
        raise MissingFileError(f"Not a file: {path}")

    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        shown = ext or "(no extension)"
        raise UnsupportedFileError(
            f"Unsupported file type '{shown}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    _check_signature(p, ext)

    if ext == ".pdf":
        raw, pages = _extract_pdf(p)
    else:
        raw, pages = _extract_docx(p), None

    text = clean_text(raw)
    if not text.strip():
        hint = " (is it a scanned PDF? OCR is not supported)" if ext == ".pdf" else ""
        raise NoTextFoundError(f"No extractable text in {p.name}{hint}")
    return ExtractedDocument(filename=p.name, text=text, pages=pages)


def clean_text(text: str) -> str:
    """Normalize extracted text so chunking and embeddings see clean input.

    * NFKC normalization (ligatures like "ﬁ" -> "fi", full-width characters, ...)
    * removes NUL bytes - PostgreSQL rejects ``\\x00`` in TEXT columns - and other
      control characters (tabs and newlines are kept)
    * re-joins words hyphenated across a line break ("cus-\\ntomer" -> "customer")
    * collapses runs of spaces, trims each line, keeps at most one blank line
    """
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C")
    text = re.sub(r"(?<=[A-Za-z])-\n(?=[a-z])", "", text)
    text = re.sub(r"[ \t ]+", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ─────────────────────────── internals ───────────────────────────

def _check_signature(p: Path, ext: str) -> None:
    try:
        with p.open("rb") as fh:
            head = fh.read(8)
    except OSError as exc:
        raise CorruptedFileError(f"Cannot read {p.name}: {exc.strerror or exc}") from exc
    if not head:
        raise NoTextFoundError(f"{p.name} is empty (0 bytes)")
    expected = PDF_MAGIC if ext == ".pdf" else ZIP_MAGIC
    if not head.startswith(expected):
        raise UnsupportedFileError(
            f"{p.name} has a '{ext}' extension but its content is not a valid {ext[1:].upper()} file"
        )


def _extract_pdf(p: Path) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(p))
        if reader.is_encrypted:
            # Many PDFs are "encrypted" with an empty user password; try it before giving up.
            try:
                ok = reader.decrypt("")
            except Exception:  # pypdf raises different types depending on the cipher
                ok = 0
            if not ok:
                raise CorruptedFileError(f"{p.name} is password-protected and cannot be read")
        pages = [_pdf_page_to_text(page.extract_text() or "") for page in reader.pages]
    except CorruptedFileError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise CorruptedFileError(f"{p.name} is corrupted or not a readable PDF ({exc})") from exc
    return "\n\n".join(pg for pg in pages if pg.strip()), len(reader.pages)


def _pdf_page_to_text(page_text: str) -> str:
    """Rebuild paragraphs from PDF lines.

    PDF text comes out one visual line at a time, with no blank lines between
    paragraphs. A line break is treated as a paragraph break when the line looks
    like the end of a block: it ends with terminal punctuation, or it is clearly
    shorter than a typical full line (a heading or the last line of a paragraph).
    A line that ends with a full stop but fills the whole width is usually a
    sentence that happens to end at the margin, so it does not end the block.
    Bullet lines always start a new block. Other line breaks become spaces.
    """
    lines = [ln.strip() for ln in page_text.replace("\r", "\n").split("\n")]
    lines = [ln for ln in lines if ln]
    if not lines:
        return ""
    lengths = sorted(len(ln) for ln in lines)
    # "typical full line": 75th percentile, so short table cells don't drag it down
    typical = lengths[int(0.75 * (len(lengths) - 1))] if len(lengths) >= 3 else 80
    bullet = re.compile(r"^([-•*·▪]|\d+[.)])\s")

    out: list[str] = []
    buf = last = lines[0]
    for ln in lines[1:]:
        short = len(last) < 0.6 * typical
        sentence_end = bool(re.search(r"[.!?]$", last)) and len(last) < 0.9 * typical
        prev_ends_block = short or sentence_end
        if prev_ends_block or bullet.match(ln):
            out.append(buf)
            buf = ln
        else:
            buf = f"{buf[:-1]}{ln}" if buf.endswith("-") and ln[:1].islower() else f"{buf} {ln}"
        last = ln
    out.append(buf)
    return "\n\n".join(out)


def _extract_docx(p: Path) -> str:
    import docx  # python-docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(str(p))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise CorruptedFileError(f"{p.name} is corrupted or not a valid Word document ({exc})") from exc

    blocks: list[str] = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:  # tables are common in procedure documents
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat; dedupe
    return "\n\n".join(blocks)
