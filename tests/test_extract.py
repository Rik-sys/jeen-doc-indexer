"""Extraction and cleaning tests - use temporary files, no API key or database needed."""

from pathlib import Path

import docx
import pytest

from doc_indexer.errors import CorruptedFileError, MissingFileError, NoTextFoundError, UnsupportedFileError
from doc_indexer.extract import _pdf_page_to_text, clean_text, extract_document

DOCS = Path(__file__).resolve().parent.parent / "docs"


class TestCleanText:
    def test_removes_nul_and_control_characters(self):
        assert clean_text("abc\x00def\x07") == "abcdef"

    def test_nfkc_normalization(self):
        assert clean_text("ﬁle") == "file"  # "fi" ligature

    def test_rejoins_hyphenated_words(self):
        assert clean_text("cus-\ntomer") == "customer"

    def test_collapses_whitespace_and_blank_lines(self):
        assert clean_text("a   b\t\tc\n\n\n\n\nd ") == "a b c\n\nd"

    def test_keeps_hebrew(self):
        assert clean_text("שלום  עולם") == "שלום עולם"


class TestPdfParagraphs:
    def test_joins_wrapped_lines_and_splits_on_headings(self):
        page = "1. Heading\nThis line wraps at the margin and the sentence continues\non the next line until it ends here.\n2. Next"
        out = _pdf_page_to_text(page).split("\n\n")
        assert out[0] == "1. Heading"
        assert out[1].startswith("This line wraps") and "continues on the next" in out[1]


class TestExtractDocument:
    def test_sample_pdf(self):
        doc = extract_document(DOCS / "example.pdf")
        assert doc.filename == "example.pdf" and doc.pages and "KB-112" in doc.text

    def test_sample_docx_hebrew(self):
        doc = extract_document(DOCS / "technician_visits_he.docx")
        assert "טכנאי" in doc.text and doc.pages is None

    def test_docx_tables_are_extracted(self, tmp_path):
        d = docx.Document()
        d.add_paragraph("Intro")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "SLA", "4 hours"
        path = tmp_path / "t.docx"
        d.save(path)
        assert "SLA | 4 hours" in extract_document(path).text

    def test_missing_file(self):
        with pytest.raises(MissingFileError):
            extract_document("does/not/exist.pdf")

    def test_unsupported_extension(self, tmp_path):
        f = tmp_path / "notes.txt"
        f.write_text("hello")
        with pytest.raises(UnsupportedFileError):
            extract_document(f)

    def test_wrong_content_for_extension(self, tmp_path):
        f = tmp_path / "fake.pdf"
        f.write_text("this is not a pdf")
        with pytest.raises(UnsupportedFileError):
            extract_document(f)

    def test_empty_file(self, tmp_path):
        f = tmp_path / "empty.pdf"
        f.write_bytes(b"")
        with pytest.raises(NoTextFoundError):
            extract_document(f)

    def test_docx_without_text(self, tmp_path):
        path = tmp_path / "blank.docx"
        docx.Document().save(path)
        with pytest.raises(NoTextFoundError):
            extract_document(path)

    def test_corrupted_pdf(self, tmp_path):
        f = tmp_path / "broken.pdf"
        f.write_bytes(b"%PDF-1.7\n" + b"\x00garbage" * 50)
        with pytest.raises((CorruptedFileError, NoTextFoundError)):
            extract_document(f)

    def test_corrupted_docx(self, tmp_path):
        f = tmp_path / "broken.docx"
        f.write_bytes(b"PK\x03\x04" + b"not really a zip" * 20)
        with pytest.raises(CorruptedFileError):
            extract_document(f)
